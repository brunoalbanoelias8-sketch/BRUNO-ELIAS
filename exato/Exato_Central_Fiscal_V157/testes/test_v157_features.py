"""V157: PDF de cada NFS-e (mesma lógica do XML da NF-e), exportação marcada, alerta de cancelada exportada,
busca de todas as empresas, resumo mensal, Excel, cartões do Início e modo técnico."""
import json, os, sys, tempfile, shutil, time, zipfile
from decimal import Decimal
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
os.environ['EXATO_DATA_DIR']=tempfile.mkdtemp(prefix='exato_v157_'); sys.path.insert(0,str(ROOT/'programa')); sys.path.insert(0,str(ROOT/'programa'/'third_party')); sys.path.insert(0,str(Path(__file__).parent))
import exato_central_fiscal as m, exato_nfse_pdf as npdf, exato_xlsx
from test_v157_nfse_core import nfse_xml, evento_xml, chave, pack, PREST, TOMA
from pypdf import PdfReader
m.init_database()
u=m.db_auth_get_user(m.AUTH_BOOTSTRAP_EMAIL); m.db_auth_set_password(u['id'],'SenhaUI#2026'); user=m.db_auth_login(m.AUTH_BOOTSTRAP_EMAIL,'SenhaUI#2026')
m.enumerate_windows_certificates=lambda *a,**k:([],'')
msgs=[]; m.messagebox.showinfo=lambda *a,**k:msgs.append(('info',)+a); m.messagebox.showwarning=lambda *a,**k:msgs.append(('warn',)+a); m.messagebox.showerror=lambda *a,**k:msgs.append(('erro',)+a); m.messagebox.askyesno=lambda *a,**k:True
out=Path(tempfile.mkdtemp(prefix='exato_v157_out_')); opened=[]; m._open_default_path=lambda p:opened.append(str(p))
def text_of(path): return ' '.join(p.extract_text() for p in PdfReader(str(path)).pages)
OTHER='11222333000181'
try:
    # ---------- notas: 2 prestadas (uma cancelada) e 1 tomada
    items=[{'nsu':1,'chave':chave(1),'tipo_documento':'NFSE','xml':nfse_xml(1,valor='1000.00',dh='2026-09-10T10:00:00-03:00')},
           {'nsu':2,'chave':chave(2),'tipo_documento':'NFSE','xml':nfse_xml(2,valor='250.50',dh='2026-09-12T10:00:00-03:00')},
           {'nsu':3,'chave':chave(3),'tipo_documento':'NFSE','xml':nfse_xml(3,prest=OTHER,toma=PREST,valor='80.00',dh='2026-10-02T10:00:00-03:00')}]
    m.db_upsert_nfse_items(PREST,items)
    docs=m.db_load_documents_as_payload(PREST,family='nfse')
    # ---------- exportação como a NF-e: XML + PDF de cada nota + relatório mensal da pasta
    saved,errors=m.save_nfse_documents(docs,str(out/'a'),PREST,generate_companion_pdf=True)
    xmls=[x for x in saved if x.endswith('.xml')]; pdfs=[x for x in saved if x.endswith('.pdf')]
    assert not errors and len(xmls)==3 and len(pdfs)==3+2,(len(xmls),len(pdfs),errors)       # 3 individuais + 2 relatórios mensais (Prestados 09/2026 e Tomados 10/2026)
    cons=[x for x in pdfs if 'Relatorio_mensal' in x]; assert len(cons)==2 and any('Prestados' in x and '2026-09' in x for x in cons) and any('Tomados' in x and '2026-10' in x for x in cons)
    prest_cons=next(x for x in cons if 'Prestados' in x); rt=text_of(prest_cons); assert 'Total do mês' in rt and 'Relatório mensal de NFS-e' in rt and 'setembro/2026' in rt and rt.count('Autorizada')>=2     # as 2 notas de setembro, com o total
    single=next(x for x in pdfs if x.endswith(chave(1)+'.pdf')); assert len(PdfReader(single).pages)==1 and 'DANFSe' in text_of(single) and 'PRESTADORA SERVICOS LTDA' in text_of(single) and '1.000,00' in text_of(single)
    # sem PDF (como "Salvar XMLs" de Documentos na NF-e): só XML
    saved2,_=m.save_documents_any(docs,str(out/'b'),PREST,{},None,run_legacy_migration=False,generate_companion_pdf=False); assert all(x.endswith('.xml') for x in saved2) and len(saved2)==3
    # ---------- cancelamento da nota 2 depois de exportada: alerta
    m.db_mark_documents_exported([d for d in docs if d['chave']!=chave(3)],str(out/'a'))
    r=m.db_upsert_nfse_items(PREST,[{'nsu':4,'chave':chave(2),'tipo_documento':'EVENTO','tipo_evento':'101101','xml':evento_xml(2)}]); assert r['cancelled']==1 and r['cancelled_exported']==1,r
    docs2={d['chave']:d for d in m.db_load_documents_as_payload(PREST,family='nfse')}; assert docs2[chave(2)]['status']=='Cancelado'
    p=m.open_nfse_representation(next(x for x in m.db_list_documents(cnpj=PREST,family='nfse') if x['access_key']==chave(2) and x['status']!='Evento')); assert opened and 'CANCELADA' in text_of(p)
    # ---------- representação PDF e totais para o Início
    tot=m.db_nfse_totals(PREST); assert (tot['prest_qtd'],tot['prest_valor'],tot['tom_qtd'],tot['canc_qtd'])==(1,Decimal('1000'),1,1),tot
    assert m.db_get_stats(PREST)['total']==0 and m.db_get_stats(PREST)['authorized']==0, 'NFS-e não entram nos totais das demais notas'
    # ---------- resumo mensal
    rows=[{'data':'2026-09-10','tipo':'Prestado','valor':'1000.00','situacao':'Autorizada'},{'data':'2026-09-12','tipo':'Prestado','valor':'250.50','situacao':'Cancelada'},{'data':'2026-10-02','tipo':'Tomado','valor':'80.00','situacao':'Autorizada'}]
    ms=npdf.monthly_summary(rows); assert [(x['mes'],x['prest_qtd'],str(x['prest_valor']),x['tom_qtd'],x['canc_qtd']) for x in ms]==[('2026-09',1,'1000.00',0,1),('2026-10',0,'0.00',1,0)]
    assert npdf.month_label('2026-09')=='setembro/2026'
    # ---------- aplicativo
    app=m.App(current_user=user); app.geometry('1366x720+0+0')
    def loop(cond,t=10,start=None):
        end=time.time()+t; res=[False]
        def tick():
            if cond(): res[0]=True; app.quit()
            elif time.time()>end: app.quit()
            else: app.after(30,tick)
        if start: app.after(0,start)
        app.after(40,tick); app.mainloop(); return res[0]
    try:
        app.update(); app._show_nfse(); app.nfse_company.set(m._format_cnpj(PREST)); app._nfse_refresh_table(); app.update()
        vals={app.nfse_tree.item(i,'values')[0]:app.nfse_tree.item(i,'values') for i in app.nfse_tree.get_children()}
        assert vals['1'][7]=='✓ Exportada' and '✕ Cancelada · já exportada' in vals['2'][6] and vals['3'][7]=='○ Não exportada', vals
        assert 'ATENÇÃO: 1 nota(s) já exportada(s) foram canceladas' in app.nfse_shown_label.cget('text')
        for kind,expected in (('Exportadas',{'1','2'}),('Não exportadas',{'3'}),('Canceladas já exportadas',{'2'})):
            app.nfse_export_filter.set(kind); app._nfse_refresh_table(); assert {app.nfse_tree.item(i,'values')[0] for i in app.nfse_tree.get_children()}==expected,kind
        app.nfse_export_filter.set('Todas'); app._nfse_refresh_table()
        # salvar com PDF pela aba (padrão ligado) e "somente novos"
        m.filedialog.askdirectory=lambda **k:str(out/'c')
        app.after(0,app._nfse_save_xmls); assert loop(lambda: len(list((out/'c').rglob('*.pdf')))>=5), 'XMLs, PDFs e consolidados'
        assert len(list((out/'c').rglob('*.xml')))==3
        assert loop(lambda: not m.db_filter_unexported(m.db_load_documents_as_payload(PREST,family='nfse'),str(out/'c'))), 'exportação registrada'
        msgs.clear(); app._nfse_save_xmls(True); app.update(); assert any('já registrados' in str(x) or 'já foram' in str(x) or 'exportados' in str(x) for x in msgs), msgs
        # PDF individual pela lista e pelo lote de Documentos
        app.nfse_tree.selection_set(app.nfse_tree.get_children()[0]); opened.clear(); app._nfse_open_selected_pdf(); assert opened and opened[0].endswith('.pdf')
        app._show_documents(); app.doc_family.set('NFS-e'); app._refresh_documents_list(); app.update()
        kids=app.doc_tree.get_children(); assert len(kids)==3
        statuses=[app.doc_tree.item(i,'values')[6] for i in kids]; assert any('Cancelado · já exportado' in s for s in statuses),statuses
        shutil.rmtree(m.APP_DATA_DIR/'Representacoes_Fiscais',ignore_errors=True)
        app.doc_tree.selection_set(kids); app.update()
        app.after(0,app._generate_selected_representations_batch); assert loop(lambda: len(list((m.APP_DATA_DIR/'Representacoes_Fiscais').rglob('NFSe_*.pdf')))==3), 'lote de representações'
        # ---------- resumo mensal e Excel pela aba
        app._show_nfse(); app.nfse_company.set(m._format_cnpj(PREST)); app._nfse_refresh_table(); app.update(); app.nfse_tree.selection_set(())
        m.filedialog.asksaveasfilename=lambda **k:str(out/('resumo.pdf' if k.get('defaultextension')=='.pdf' else 'planilha.xlsx'))
        app._nfse_monthly_pdf(); t=text_of(out/'resumo.pdf'); assert 'Resumo mensal de NFS-e' in t and 'setembro/2026' in t and 'outubro/2026' in t and 'Canceladas' in t
        app._nfse_monthly_excel(); app.update()
        with zipfile.ZipFile(out/'planilha.xlsx') as zf: assert 'xl/worksheets/sheet2.xml' in zf.namelist() and b'setembro/2026' in zf.read('xl/worksheets/sheet1.xml')
        try:
            import openpyxl
            wb=openpyxl.load_workbook(out/'planilha.xlsx'); assert wb.sheetnames==['Resumo mensal','Notas']
            assert [c.value for c in wb['Resumo mensal'][2]][:3]==['setembro/2026',1,1000], [c.value for c in wb['Resumo mensal'][2]]
            assert len(list(wb['Notas'].iter_rows()))==4 and wb['Notas']['F2'].number_format=='#,##0.00'
            app._nfse_list_excel(); wb=openpyxl.load_workbook(out/'planilha.xlsx'); assert wb.sheetnames==['Notas']
        except ImportError: pass
        # ---------- Início: cartões de NFS-e
        app.cnpj_var.set(PREST); app._show_dashboard(); app.update(); app._apply_dashboard_stats(m.db_get_stats(PREST),'EMPRESA',[],[],[],{},m.db_nfse_totals(PREST)); app.update()
        assert '1.000,00' in app.dash_values['nfse_prest'].cget('text') and '1 nota' in app.dash_values['nfse_prest'].cget('text') and app.dash_values['nfse_canc'].cget('text')=='1' and '80,00' in app.dash_values['nfse_tom'].cget('text')
        # ---------- modo técnico escondido por padrão
        app._show_maintenance(); app.update(); assert not app.maint_tech_var.get() and not app.maint_tech_frame.winfo_manager()
        app.maint_tech_var.set(True); app._apply_technical_mode(); app.update(); assert app.maint_tech_frame.winfo_manager() and app.config_data['technical_mode'] is True
        app.maint_tech_var.set(False); app._apply_technical_mode(); app.update(); assert not app.maint_tech_frame.winfo_manager()
        # ---------- buscar todas as empresas
        CN1,CN2,CN3,CN4='12345678000195','11222333000181','34028316000103','00000000000191'
        for c,n in ((CN1,'ALFA LTDA'),(CN2,'BETA LTDA'),(CN3,'GAMA LTDA'),(CN4,'SEM CERTIFICADO LTDA')): m.db_register_company(c,n)
        app.certificates=[{'Thumbprint':'T1','FriendlyName':'ALFA','Document':m._format_cnpj(CN1),'NotAfter':'2099-01-01'},{'Thumbprint':'T2','FriendlyName':'BETA','Document':m._format_cnpj(CN2),'NotAfter':'2099-01-01'},{'Thumbprint':'T3','FriendlyName':'GAMA','Document':m._format_cnpj(CN3),'NotAfter':'2099-01-01'}]
        pairs,missing=app._nfse_company_certificates(); assert sorted(c['cnpj'] for c,_ in pairs)==sorted([CN1,CN2,CN3]) and [c['cnpj'] for c in missing].count(CN4)==1
        calls=[]
        def fake(url,thumb):
            calls.append((url,thumb)); cn=url.split('cnpjConsulta=')[1]; nsu=int(url.split('/DFe/')[1].split('?')[0])
            if cn==CN3: return (403,'{}')
            if nsu>0: return (404,json.dumps({'StatusProcessamento':'NENHUM_DOCUMENTO_LOCALIZADO'}))
            xml=nfse_xml(50 if cn==CN1 else 60,prest=cn,toma='98765432000110')
            return (200,json.dumps({'LoteDFe':[{'NSU':1,'ChaveAcesso':chave(50 if cn==CN1 else 60),'TipoDocumento':'NFSE','ArquivoXml':pack(xml)}]}))
        app._nfse_transport=fake; m.nfse_mod.ADN_PAUSE_SECONDS=0
        app._show_nfse(); assert loop(lambda: not app._nfse_running, start=app._nfse_run_all) and calls
        got=[m.db_nfse_stats(c)['total'] for c in (CN1,CN2,CN3,CN4)]; assert got==[4,1,0,0],(got,app.nfse_status.cget('text'),app.nfse_detail.cget('text'),calls[:6],[(r['cnpj'],r['status'],r['error_text']) for r in m.db_list_runs(6)])
        assert 'com problema' in app.nfse_status.cget('text') and 'GAMA LTDA' in app.nfse_detail.cget('text') and '2 empresa(s) consultada(s)' in app.nfse_detail.cget('text'),(app.nfse_status.cget('text'),app.nfse_detail.cget('text'))
        assert {t for _,t in calls}=={'T1','T2','T3'}
        assert m.db_list_runs(1,CN3)[0]['status']=='Não concluído' and m.db_list_runs(1,CN1)[0]['status']=='Concluído'
    finally:
        app.destroy()
finally:
    shutil.rmtree(os.environ['EXATO_DATA_DIR'],ignore_errors=True); shutil.rmtree(out,ignore_errors=True)
print('V157 features: OK')
